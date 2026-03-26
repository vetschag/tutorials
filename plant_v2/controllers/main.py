import io
import base64
import logging

from odoo import http
from odoo.http import request, Response

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_logger = logging.getLogger(__name__)


def set_heading_color(paragraph, r, g, b):
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(r, g, b)


def add_bottom_border(paragraph, color="1F3864"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def decode_image(img_raw):
    """Konvertiert Odoo Binary Feld zu bytes"""
    if not img_raw:
        return None
    try:
        if isinstance(img_raw, bytes):
            # Prüfen ob es base64-kodierte Bytes sind
            try:
                decoded = base64.b64decode(img_raw)
                _logger.info("Image: bytes → base64 decode → %d bytes", len(decoded))
                return decoded
            except Exception:
                _logger.info("Image: raw bytes → %d bytes", len(img_raw))
                return img_raw
        elif isinstance(img_raw, str):
            decoded = base64.b64decode(img_raw)
            _logger.info("Image: string → base64 decode → %d bytes", len(decoded))
            return decoded
        else:
            _logger.warning("Image: unbekannter Typ: %s", type(img_raw))
            return None
    except Exception as e:
        _logger.error("Image decode Fehler: %s", e)
        return None


class PlantTemplateDocxController(http.Controller):

    @http.route('/plant/template/<int:template_id>/docx',
                type='http', auth='user', methods=['GET'])
    def download_docx(self, template_id, **kwargs):

        template = request.env['plant.template'].browse(template_id)
        if not template.exists():
            return Response("Template nicht gefunden", status=404)

        doc = Document()

        # Seitenränder
        section = doc.sections[0]
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

        doc.styles['Normal'].font.name = 'Arial'
        doc.styles['Normal'].font.size = Pt(11)

        # Titel
        title = doc.add_heading(template.name, level=0)
        set_heading_color(title, 0x1F, 0x38, 0x64)
        title.runs[0].font.size = Pt(20)

        p = doc.add_paragraph()
        run = p.add_run(f"Plant: {template.plant_id.name or ''}")
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        doc.add_page_break()

        # Inhaltsverzeichnis
        toc = doc.add_paragraph()
        toc.add_run("Inhaltsverzeichnis").bold = True
        toc.runs[0].font.size = Pt(13)

        block_num = 0
        for block in template.plant_id.block_ids:
            block_num += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            run = p.add_run(f"{block_num}.  {block.name}")
            run.bold = True

            agg_num = 0
            for agg in block.aggregate_ids:
                agg_num += 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.5)
                p.add_run(f"{block_num}.{agg_num}.  {agg.name}").font.size = Pt(10)

        doc.add_page_break()

        # Inhalt
        block_num = 0
        for block in template.plant_id.block_ids:
            block_num += 1

            h1 = doc.add_heading(f"{block_num}. {block.name}", level=1)
            set_heading_color(h1, 0x1F, 0x38, 0x64)
            add_bottom_border(h1)

            agg_num = 0
            for agg in block.aggregate_ids:
                agg_num += 1

                h2 = doc.add_heading(f"{block_num}.{agg_num}. {agg.name}", level=2)
                set_heading_color(h2, 0x2E, 0x75, 0xB6)

                # Text + Selection als Bullet Points
                for type_ in agg.type_ids:
                    for field in type_.field_ids:
                        value = template.value_ids.filtered(
                            lambda v, fid=field.id: v.field_id.id == fid
                        )
                        val = value[0] if value else None

                        if field.field_type == 'text' and val and val.value_text:
                            p = doc.add_paragraph(style='List Bullet')
                            p.add_run(f"{field.name}: ").bold = True
                            p.add_run(val.value_text)

                        elif field.field_type == 'selection' and val and val.value_selection_id:
                            p = doc.add_paragraph(style='List Bullet')
                            p.add_run(f"{field.name}: ").bold = True
                            p.add_run(val.value_selection_id.name)

                # Bilder
                for type_ in agg.type_ids:
                    for field in type_.field_ids:
                        if field.field_type == 'image':
                            value = template.value_ids.filtered(
                                lambda v, fid=field.id: v.field_id.id == fid
                            )
                            val = value[0] if value else None
                            if val and val.value_image:
                                try:
                                    # Odoo gibt Binary immer als base64 zurück
                                    img_raw = val.with_context(bin_size=False).value_image
                                    if isinstance(img_raw, str):
                                        img_bytes = base64.b64decode(img_raw)
                                    elif isinstance(img_raw, bytes):
                                        # Prüfen ob es base64-String als bytes ist
                                        try:
                                            img_bytes = base64.b64decode(img_raw)
                                        except Exception:
                                            img_bytes = img_raw
                                    else:
                                        continue

                                    img_stream = io.BytesIO(img_bytes)
                                    img_stream.seek(0)
                                    p = doc.add_paragraph()
                                    p.add_run().add_picture(img_stream, width=Cm(8))
                                    _logger.info("Bild OK: %s (%d bytes)", field.name, len(img_bytes))
                                except Exception as e:
                                    _logger.error("Bild Fehler %s: %s", field.name, e)

                doc.add_page_break()

        # Speichern
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = f"{template.name.replace(' ', '_')}.docx"
        return Response(
            buffer.read(),
            headers={
                'Content-Type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'Content-Disposition': f'attachment; filename="{filename}"',
            }
        )
